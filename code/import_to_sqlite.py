import pandas as pd
import sqlite3
import argparse
import logging
from pathlib import Path
from datetime import datetime
from docx import Document
from docx.shared import Inches
import plotly.express as px

# === Set up logging ===
logging.basicConfig(
    level=logging.INFO, format="%(asctime)s - %(levelname)s - %(message)s"
)

# === Parse arguments ===
parser = argparse.ArgumentParser(
    description="Import emissions data and generate compliance report."
)
parser.add_argument(
    "--report_date",
    type=str,
    default=datetime.now().strftime("%Y-%m-%d"),
    help="Report date (YYYY-MM-DD)",
)
args = parser.parse_args()

# === Set up paths ===
project_root = Path(__file__).resolve().parent.parent
output_dir = project_root / "output"
db_path = project_root / "data" / "emissions.db"
report_path = project_root / "reports"
daily_report_folder = report_path / args.report_date
daily_report_folder.mkdir(parents=True, exist_ok=True)
summary_csv_path = daily_report_folder / f"summary_{args.report_date}.csv"
db_path.parent.mkdir(parents=True, exist_ok=True)
report_path.mkdir(parents=True, exist_ok=True)

# === Connect to SQLite DB ===
conn = sqlite3.connect(db_path)
cursor = conn.cursor()

cursor.execute("DROP TABLE IF EXISTS flights")

# === Create table ===
cursor.execute(
    """
CREATE TABLE IF NOT EXISTS flights (
    flight_id TEXT,
    date TEXT,
    origin TEXT,
    destination TEXT,
    aircraft_type TEXT,
    flight_duration_hr REAL,
    fuel_burn_kg REAL,
    SAF_blend_pct REAL,
    co2_emissions_kg REAL,
    origin_country TEXT,
    destination_country TEXT,
    in_eea INTEGER,
    operator_id TEXT,
    operator_name TEXT,
    operator_compliant INTEGER,
    verified INTEGER,
    offset_required INTEGER,
    imported_at TEXT
)
"""
)

cursor.execute("CREATE INDEX IF NOT EXISTS idx_flights_date ON flights(date);")
cursor.execute(
    "CREATE INDEX IF NOT EXISTS idx_operator_name ON flights(operator_name);"
)
cursor.execute("CREATE INDEX IF NOT EXISTS idx_imported_at ON flights(imported_at);")

# === Track previously imported files ===
imported_files_path = output_dir / ".imported_files.txt"
if imported_files_path.exists():
    logging.info("🗑️ Deleting outdated .imported_files.txt to ensure fresh import")
    imported_files_path.unlink()
imported = set()

# === Import new CSVs ===
csv_files = sorted(output_dir.glob("flights_202*.csv"))
new_rows = 0
for csv_file in csv_files:
    if csv_file.name in imported:
        continue

    logging.info(f"Importing {csv_file.name}")
    df = pd.read_csv(csv_file)
    df["date"] = pd.to_datetime(df["date"]).dt.strftime("%Y-%m-%d")
    df["imported_at"] = datetime.now().isoformat(timespec="seconds")
    df.to_sql("flights", conn, if_exists="append", index=False)
    imported.add(csv_file.name)
    new_rows += len(df)

imported_files_path.write_text("\n".join(imported))

# === Generate summary metrics ===
df_summary = pd.read_sql_query(
    """
    SELECT 
        COUNT(*) as total_flights,
        ROUND(SUM(co2_emissions_kg) / 1000.0, 2) as total_emissions_tonnes,
        ROUND(AVG(in_eea) * 100, 2) as eea_percent,
        ROUND(AVG(verified) * 100, 2) as verified_percent,
        ROUND(AVG(operator_compliant) * 100, 2) as compliant_percent
    FROM flights
""",
    conn,
)

# === Save CSV Summary ===
df_summary.to_csv(summary_csv_path, index=False)
logging.info(f"📊 Summary CSV exported to {summary_csv_path}")

# === Append to master summary log without duplicates ===
summary_log_path = report_path / "summary_log.csv"
df_summary["report_date"] = args.report_date

if summary_log_path.exists():
    df_log = pd.read_csv(summary_log_path)
    if args.report_date not in df_log["report_date"].astype(str).values:
        df_log = pd.concat([df_log, df_summary], ignore_index=True)
        df_log.to_csv(summary_log_path, index=False)
        logging.info(f"📝 Summary appended to {summary_log_path}")
    else:
        logging.info(
            f"⚠️ Summary for {args.report_date} already exists in log. Skipping append."
        )
else:
    df_summary.to_csv(summary_log_path, index=False)
    logging.info(f"📘 Created new summary log at {summary_log_path}")
logging.info(f"📊 Summary CSV exported to {summary_csv_path}")

# === Create Word report ===
report = Document()
report.add_heading("Emissions Compliance Summary Report", 0)
report.add_paragraph(f"Generated on {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")

# === Add KPI table ===
table = report.add_table(rows=1, cols=2)
table.style = "Light List Accent 1"
hdr = table.rows[0].cells
hdr[0].text = "Metric"
hdr[1].text = "Value"
for col in df_summary.columns:
    row = table.add_row().cells
    row[0].text = col.replace("_", " ").title()
    row[1].text = str(df_summary[col][0])

# === Prepare data for plots ===
df_plot = pd.read_sql_query("SELECT * FROM flights", conn)
df_plot["in_eea"] = df_plot["in_eea"].astype(bool)
df_plot["operator_compliant"] = df_plot["operator_compliant"].astype(bool)

logging.info("\n🔍 Troubleshooting Data Check")
logging.info(f"Data shape: {df_plot.shape}")
logging.info(f"Unique aircraft types: {df_plot['aircraft_type'].unique()}")
logging.info(f"Total emissions: {df_plot['co2_emissions_kg'].sum()}")

# === Create date-based report folder ===
daily_report_folder = report_path / args.report_date
daily_report_folder.mkdir(parents=True, exist_ok=True)

# === Plot 1: Emissions by aircraft ===
emissions_by_aircraft = (
    df_plot.groupby("aircraft_type")["co2_emissions_kg"].sum().reset_index()
)
if (
    not emissions_by_aircraft.empty
    and emissions_by_aircraft["co2_emissions_kg"].sum() > 0
):
    fig1 = px.bar(
        emissions_by_aircraft,
        x="aircraft_type",
        y="co2_emissions_kg",
        title="CO₂ Emissions by Aircraft Type",
    )
else:
    fig1 = px.bar(
        x=["No emissions data"], y=[0], title="CO₂ Emissions by Aircraft Type"
    )
fig1_path = daily_report_folder / "emissions_by_aircraft.png"
try:
    fig1.write_image(str(fig1_path))
except Exception as e:
    logging.error(f"Failed to save bar plot: {e}")

# === Plot 2: Compliance status ===
if df_plot["operator_compliant"].nunique() > 1:
    fig2 = px.pie(
        df_plot,
        names="operator_compliant",
        title="Compliance Status Distribution",
        hole=0.4,
        labels={True: "Compliant", False: "Non-Compliant"},
    )
else:
    fig2 = px.pie(
        names=["Only one category"],
        values=[1],
        title="Compliance Status Distribution",
        hole=0.4,
    )
fig2_path = daily_report_folder / "compliance_pie.png"
try:
    fig2.write_image(str(fig2_path))
except Exception as e:
    logging.error(f"Failed to save pie chart: {e}")

# === Insert plots ===
report.add_heading("Visual Summary", level=1)
report.add_paragraph(
    "The chart below shows the total CO₂ emissions per aircraft type. "
    "This helps identify which aircraft models contribute most to total emissions."
)
report.add_paragraph("CO₂ Emissions by Aircraft Type")
report.add_picture(str(fig1_path), width=Inches(5.5))

report.add_paragraph(
    "The following pie chart illustrates the proportion of flights that were deemed compliant "
    "under EU Emissions Trading Scheme (EU ETS) criteria. Compliance indicates that the flight's "
    "emissions data has been verified and the operator has met regulatory obligations. "
    "Non-compliant flights may reflect reporting gaps, verification issues, or allowance shortfalls."
)
report.add_paragraph("Compliance Status Distribution")
report.add_picture(str(fig2_path), width=Inches(5.5))

# === Save Word report ===
report_file = daily_report_folder / "summary.docx"
report.save(report_file)

conn.close()

logging.info(f"\n✅ {new_rows} new rows imported.")
logging.info(f"📄 Word summary report saved to: {report_file}")
